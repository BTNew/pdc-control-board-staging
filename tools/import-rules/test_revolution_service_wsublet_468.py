"""r468 (F81): the w/Sublet report with Revolution's optional leading Type column (D data, T1 footer)."""
import hashlib
import json
import os
import tempfile
import unittest

import fixtures_service_wsublet_468 as F
from revolution_service_wsublet import (TYPE_ABSENT, TYPE_BASIS, TYPE_DROPPED, ContractError, check_headers,
                                        classify_typed, collapse, has_type_column, read_values, source_type_column)

# sha256 of read_values() + collapse() on the no-Type fixture, computed with the r442/r444 module (origin/main 88602a4b)
# before this change: a file without Type must still parse exactly as today.
NO_TYPE_DIGEST_BEFORE_R468 = '2da7730703297580a4863df3005fcf888b6c6a66ee18d592c5b4d8c76dbe70dc'


def parse(values):
    h, data, footer, excluded, params = read_values(values)
    lines, summary = collapse(data)
    return h, data, footer, excluded, params, lines, summary


def digest(parsed):
    return hashlib.sha256(json.dumps(list(parsed), sort_keys=True, default=str).encode()).hexdigest()


class TypeColumn468(unittest.TestCase):
    def test_without_type_parses_exactly_as_before(self):
        self.assertEqual(digest(parse(F.build())), NO_TYPE_DIGEST_BEFORE_R468)

    def test_expected_counts_with_and_without_type(self):
        for typed in (False, True):
            h, data, footer, excluded, _, lines, s = parse(F.build(typed=typed))
            self.assertEqual(s['lines'], F.EXPECTED['lines'], typed)
            self.assertEqual(len({l['raw']['Stock #'] for l in lines}), F.EXPECTED['stocks'], typed)
            self.assertEqual(s['distinct_pos'], F.EXPECTED['distinct_pos'], typed)
            self.assertEqual((len(footer), len(excluded)), (1, 1), typed)
            self.assertGreater(s['rows_collapsed'], 0)
            self.assertEqual(s['conflicts'], [])

    def test_type_is_dropped_and_result_is_identical(self):
        plain, typed = parse(F.build()), parse(F.build(typed=True))
        self.assertEqual(digest(plain), digest(typed))
        h, data, footer, excluded, _, lines, _ = typed
        self.assertNotIn('Type', h)
        for _, raw in data + footer + excluded:
            self.assertNotIn('Type', raw)
        for l in lines:
            self.assertNotIn('Type', l['raw'])
            for c in l['sublet_columns']:
                self.assertNotIn('Type', c)

    def test_source_row_numbers_are_the_worksheet_rows(self):
        _, data, footer, _, _, lines, _ = parse(F.build(typed=True))
        self.assertEqual(data[0][0], 5)
        self.assertEqual(footer[0][0], len(F.build(typed=True)))
        self.assertEqual(lines[0]['source_row_numbers'][0], 5)

    def test_bad_type_values_stop_the_file_clearly(self):
        for bad in ('X', '', None, 'd', 'T2', 'D1', 't1', 1, 'Data'):
            with self.assertRaises(ContractError) as cm:
                read_values(F.build(typed=True, bad_type=bad, bad_at=500))
            msg = str(cm.exception)
            self.assertIn('row 505', msg, bad)
            self.assertIn('only D and T1 are accepted', msg, bad)

    def test_type_must_agree_with_the_row_shape(self):
        vals = F.build(typed=True)
        swapped = list(vals)
        swapped[-1] = ('D',) + tuple(vals[-1][1:])      # footer marked as data
        with self.assertRaisesRegex(ContractError, 'Type D .* no R/O # and no Line #'):
            read_values(swapped)
        swapped = list(vals)
        swapped[10] = ('T1',) + tuple(vals[10][1:])     # data row marked as footer
        with self.assertRaisesRegex(ContractError, 'Type T1 .* has an R/O # or Line #'):
            read_values(swapped)
        two = list(vals) + [vals[-1]]                    # a second T1 footer
        with self.assertRaisesRegex(ContractError, 'more than one footer'):
            read_values(two)

    def test_surrounding_spaces_on_type_are_trimmed(self):
        vals = list(F.build(typed=True))
        vals[7] = (' D ',) + tuple(vals[7][1:])
        vals[-1] = ('T1 ',) + tuple(vals[-1][1:])
        self.assertEqual(digest(parse(vals)), NO_TYPE_DIGEST_BEFORE_R468)

    def test_headers_and_markers(self):
        H = F.HEADERS
        self.assertTrue(has_type_column(['Type'] + H))
        self.assertTrue(has_type_column([' Type '] + H))
        self.assertFalse(has_type_column(H))
        self.assertEqual(check_headers(['Type'] + H + [None]), H)
        for bad in (H + ['Type'], H[:5] + ['Type'] + H[5:], ['Type', 'Type'] + H, ['type'] + H):
            with self.assertRaises(ContractError):
                check_headers(bad)
        self.assertEqual(source_type_column(['Type'] + H), TYPE_DROPPED)
        self.assertEqual(source_type_column(H), TYPE_ABSENT)
        self.assertEqual(TYPE_BASIS, 'absent_r442', 'the manifest marker is unchanged: no Type reaches the database')

    def test_classify_typed_department_rows(self):
        self.assertEqual(classify_typed('D', {'Dept': 135, 'R/O #': 'J135000001', 'Line #': 1}), 'excluded_department')
        self.assertEqual(classify_typed('T1', {'R/O #': None, 'Line #': None, 'Estimated labour hours': 5}), 'footer')

    def test_workbook_round_trip(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        with tempfile.TemporaryDirectory() as d:
            got = {}
            for name, kw in (('plain', {}), ('typed', {'typed': True})):
                p = F.write_xlsx(os.path.join(d, name + '.xlsx'), **kw)
                wb = openpyxl.load_workbook(p, read_only=True)
                try:
                    got[name] = digest(parse([tuple(r) for r in wb.active.iter_rows(values_only=True)]))
                finally:
                    wb.close()
            self.assertEqual(got['plain'], got['typed'])
            p = F.write_xlsx(os.path.join(d, 'bad.xlsx'), typed=True, bad_type='X', bad_at=3)
            wb = openpyxl.load_workbook(p, read_only=True)
            try:
                with self.assertRaisesRegex(ContractError, "row 8: Type 'X' is not D"):
                    read_values([tuple(r) for r in wb.active.iter_rows(values_only=True)])
            finally:
                # Read-only workbooks own an open ZIP handle even after all rows are consumed.
                wb.close()


if __name__ == '__main__':
    unittest.main()
